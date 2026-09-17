#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
HCP银行开户行智能清洗与质检引擎 (Bank Information Cleaner & Auditor) - 工业级增强版
核心解决 HCP 手填开户行疑难杂症：
1. 重复字样清洗（消除各类粘连、括号、斜杠、股份有限公司分割的重复银行名）
2. 银行简称规范化升级（如“建行太原住房支行” -> “中国建设银行太原住房支行”）
3. 缺少银行只有支行时：
   - 优先通过卡号前6位(BIN码，3177+条银联库及云端接口)自动反查并补全总行
   - 针对专属行业网点（如“住房支行”自动识别为建设银行）智能启发推断
4. 缺少支行/营业室/信用社等具体网点层级的预警与质检
5. 结尾漏字修补（如“太原解放路支” -> “太原解放路支行”）
6. 支持将清洗后的规范名称直接回填替换原开户行列，并保留原始手填留档备份
"""

import os
import re
import sys
from typing import Dict, List, Tuple, Optional
from pathlib import Path
import pandas as pd

# 适配 Windows 控制台输出 UTF-8
if hasattr(sys.stdout, 'reconfigure'):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        pass

# ==============================================================================
# 1. 全国主流银行卡 BIN 码库 (离线3177+条 + 在线兜底)
# ==============================================================================
BANK_BIN_MAP = {
    # 常用大行保底字典
    "622202": "中国工商银行", "622200": "中国工商银行", "621226": "中国工商银行",
    "622700": "中国建设银行", "621700": "中国建设银行", "436742": "中国建设银行",
    "622848": "中国农业银行", "622845": "中国农业银行", "622846": "中国农业银行",
    "621661": "中国银行", "621660": "中国银行", "601382": "中国银行",
    "622260": "交通银行", "622262": "交通银行", "601428": "交通银行",
    "621098": "中国邮政储蓄银行", "622188": "中国邮政储蓄银行",
    "621483": "招商银行", "621485": "招商银行", "622588": "招商银行",
    "622521": "上海浦东发展银行", "621789": "上海浦东发展银行",
    "622690": "中信银行", "622691": "中信银行",
    "622660": "中国光大银行", "622661": "中国光大银行",
    "622630": "华夏银行", "621222": "华夏银行",
    "622622": "中国民生银行", "622615": "中国民生银行",
    "622556": "广发银行", "622155": "平安银行",
    "622908": "兴业银行", "622365": "浙商银行",
    "622812": "北京银行", "622268": "上海银行",
    "622881": "山西省农村信用社", "622880": "晋商银行"
}

# 动态加载 3177 条全量全国卡 BIN 码库
FULL_BIN_MAP = dict(BANK_BIN_MAP)
_bin_file = os.path.join(os.path.dirname(__file__), "card_bin.json")
if os.path.exists(_bin_file):
    try:
        import json
        with open(_bin_file, "r", encoding="utf-8") as _f:
            _loaded = json.load(_f)
            FULL_BIN_MAP.update(_loaded)
    except Exception:
        pass

ALIPAY_BANK_CODE_MAP = {
    "ICBC": "中国工商银行", "ABC": "中国农业银行", "BOC": "中国银行", "CCB": "中国建设银行",
    "COMM": "交通银行", "PSBC": "中国邮政储蓄银行", "CMB": "招商银行", "SPDB": "上海浦东发展银行",
    "CITIC": "中信银行", "CEB": "中国光大银行", "HXBANK": "华夏银行", "CMBC": "中国民生银行",
    "GDB": "广发银行", "SPABANK": "平安银行", "CIB": "兴业银行", "BOSH": "上海银行",
    "BJBANK": "北京银行", "CZBANK": "浙商银行", "CBHB": "渤海银行", "EGBANK": "恒丰银行",
    "JSBANK": "江苏银行", "NJCB": "南京银行", "NBBANK": "宁波银行", "HZCB": "杭州银行",
    "WZCB": "温州银行", "CSCB": "长沙银行", "CQBANK": "重庆银行", "CDCB": "成都银行"
}

_ONLINE_BIN_CACHE: Dict[str, str] = {}

# 合法网点层级关键词
BRANCH_KEYWORDS = [
    "支行", "分行", "营业部", "营业室", "分理处",
    "信用社", "农村信用社", "合作社", "储蓄所", "村镇银行"
]

# 常见核心银行与别名
KNOWN_BANK_KEYWORDS = [
    ("中国工商银行", ["工商银行", "工行"]),
    ("中国建设银行", ["建设银行", "建行"]),
    ("中国农业银行", ["农业银行", "农行"]),
    ("中国银行", ["中行"]),
    ("交通银行", ["交行"]),
    ("中国邮政储蓄银行", ["邮政储蓄银行", "邮储银行", "邮政储蓄", "邮政银行", "邮储"]),
    ("招商银行", ["招行"]),
    ("上海浦东发展银行", ["浦东发展银行", "浦发银行", "浦发"]),
    ("中信银行", ["中信"]),
    ("中国光大银行", ["光大银行", "光大"]),
    ("华夏银行", ["华夏"]),
    ("中国民生银行", ["民生银行", "民生"]),
    ("广发银行", ["广东发展银行"]),
    ("平安银行", ["平安"]),
    ("兴业银行", ["兴业"]),
    ("浙商银行", ["浙商"]),
    ("渤海银行", ["渤海"]),
    ("恒丰银行", ["恒丰"]),
    ("北京银行", []), ("上海银行", []), ("江苏银行", []),
    ("南京银行", []), ("宁波银行", []), ("杭州银行", []),
    ("晋商银行", []), ("农村信用社", ["农信社", "信用社", "农村合作银行", "农商行", "农村商业银行"])
]


# ==============================================================================
# 2. 核心清洗与质检算法
# ==============================================================================

def clean_card_number(val) -> str:
    """强力清洗银行卡号，防止 Excel 科学计数法或浮点数导致失真"""
    if val is None or pd.isna(val):
        return ""
    if isinstance(val, (int, float)):
        try:
            return f"{int(val)}"
        except Exception:
            pass
    s = str(val).strip()
    if s.endswith(".0") and len(s) > 2 and s[:-2].isdigit():
        s = s[:-2]
    if 'e+' in s.lower():
        try:
            s = f"{int(float(s))}"
        except Exception:
            pass
    return re.sub(r'[\s\-]+', '', s)


def identify_bank_from_card(card_no: str) -> Optional[str]:
    """根据银行卡号 (BIN 码) 智能识别所属银行"""
    clean_card = clean_card_number(card_no)
    if len(clean_card) < 3:
        return None
    
    if clean_card in _ONLINE_BIN_CACHE:
        return _ONLINE_BIN_CACHE[clean_card]

    # 从长到短在本地 3177 条全量库中扫描匹配 (8位到3位)
    for prefix_len in (8, 7, 6, 5, 4, 3):
        if len(clean_card) >= prefix_len:
            prefix = clean_card[:prefix_len]
            if prefix in FULL_BIN_MAP:
                return FULL_BIN_MAP[prefix]

    # 在线官方云端接口兜底 (限卡号>=15位且未命中本地库)
    if len(clean_card) >= 15:
        try:
            import urllib.request
            import json
            api_url = f"https://ccdcapi.alipay.com/validateAndCacheCardInfo.json?_input_charset=utf-8&cardNo={clean_card}&cardBinCheck=true"
            req = urllib.request.Request(api_url, headers={"User-Agent": "Antigravity/1.0"})
            with urllib.request.urlopen(req, timeout=1.5) as resp:
                data = json.loads(resp.read().decode("utf-8"))
                if data.get("validated") and data.get("bank"):
                    b_code = data["bank"]
                    b_name = ALIPAY_BANK_CODE_MAP.get(b_code, b_code)
                    _ONLINE_BIN_CACHE[clean_card] = b_name
                    return b_name
        except Exception:
            pass

    return None


def advanced_dedup_and_clean_bank(text: str) -> Tuple[str, List[str]]:
    """
    终极全能开户行去重与规范化清洗引擎：
    1. 清理换行、零宽字符、BOM头及引导前缀；
    2. 处理括号包裹的重复（例：中国建设银行(建设银行太原住房支行)）；
    3. 处理斜杠、空格分隔的重复；
    4. 处理“股份有限公司”夹杂的重复；
    5. 处理“中国XX银行 + XX银行”高频重复（例：中国建设银行建设银行... / 中国民生银行民生银行...）；
    6. 处理连续相同银行词叠字；
    7. 处理全称 + 简称粘连（例：中国农业银行农行...）；
    8. 规范化开头简称（例：建行太原住房支行 -> 中国建设银行太原住房支行）；
    9. 末尾漏字修补（例：太原解放路支 -> 太原解放路支行）。
    """
    if not text:
        return "", []
    s = str(text).strip()
    actions = []

    # 1. 基础清理
    s = re.sub(r'[\r\n\t\u200b\uFEFF]+', '', s)
    prefix_pat = r'^(?:开户行|开户银行|开户网点|银行名称|支行名称|结算行|银行|网点)[：:\s]+'
    if re.search(prefix_pat, s):
        s = re.sub(prefix_pat, '', s).strip()
        actions.append('去除引导前缀')

    # 2. 括号嵌套重复
    # 例：中国建设银行(建设银行太原住房支行) -> 中国建设银行太原住房支行
    bracket_m = re.search(r'([^\s\(\)（）]+?银行)[\(（](.+?)[\)）]', s)
    if bracket_m:
        outer_b = bracket_m.group(1)
        inner_content = bracket_m.group(2)
        s = outer_b + inner_content
        actions.append('展开括号嵌套')

    # 3. 斜杠、减号、空格分隔的重复：中国建设银行/建设银行太原住房支行
    s = re.sub(r'([^\s/\\-]+?银行)[\s/\\-]+([^\s/\\-]+?银行)', r'\1\2', s)

    # 4. 股份有限公司夹杂在重复银行名之间
    s = re.sub(r'(中国[^\s]+?银行)股份有限公司([^\s]+?银行)', r'\1\2', s)

    # 5. 核心：主行 + 重复银行名清洗
    # (A) 中国XX银行 + XX银行（支持有无空格、标点）
    zg_m = re.search(r'(中国[^\s]+?银行)\s*(?:中国)?([^\s]+?银行)', s)
    if zg_m:
        p1, p2 = zg_m.group(1), zg_m.group(2)
        if p2 in p1 or p1 in p2:
            s = s[:zg_m.start()] + p1 + s[zg_m.end():]
            actions.append(f'剔除重复总行[{p2}]')

    # (B) 连续完全重复的银行词：XX银行XX银行
    dup_m = re.search(r'([^\s]+?银行)\s*\1', s)
    if dup_m:
        s = re.sub(r'([^\s]+?银行)\s*\1', r'\1', s)
        actions.append('剔除连续相同银行名')

    # (C) 全称 + 简称紧邻
    short_map = [
        ('中国农业银行', '农行'), ('中国工商银行', '工行'),
        ('中国建设银行', '建行'), ('中国银行', '中行'),
        ('交通银行', '交行'), ('招商银行', '招行'),
        ('中国邮政储蓄银行', '邮储银行'), ('中国邮政储蓄银行', '邮储'),
        ('中国民生银行', '民生银行'), ('中国光大银行', '光大银行'),
        ('上海浦东发展银行', '浦发银行'), ('广发银行', '广发'),
        ('平安银行', '平安'), ('兴业银行', '兴业')
    ]
    for full_n, short_n in short_map:
        pat = f'({full_n})\\s*{short_n}'
        if re.search(pat, s):
            s = re.sub(pat, r'\1', s)
            actions.append(f'剔除简称[{short_n}]')

    # 6. 开头简写规范化升级（例如“建行太原住房支行” -> “中国建设银行太原住房支行”，“建设银行太原康乐街支行” -> “中国建设银行太原康乐街支行”）
    prefix_standardize = [
        (r'^(?:中国)?建设银行', '中国建设银行'),
        (r'^(?:中国)?工商银行', '中国工商银行'),
        (r'^(?:中国)?农业银行', '中国农业银行'),
        (r'^(?:中国)?邮政储蓄银行?', '中国邮政储蓄银行'),
        (r'^(?:中国)?民生银行?', '中国民生银行'),
        (r'^(?:中国)?光大银行?', '中国光大银行'),
        (r'^(?:上海)?浦东发展银行?', '上海浦东发展银行'),
        (r'^建行', '中国建设银行'),
        (r'^工行', '中国工商银行'),
        (r'^农行', '中国农业银行'),
        (r'^中行', '中国银行'),
        (r'^交行', '交通银行'),
        (r'^邮储银行?', '中国邮政储蓄银行'),
        (r'^邮储', '中国邮政储蓄银行'),
        (r'^邮政银行', '中国邮政储蓄银行'),
        (r'^招行', '招商银行'),
        (r'^民生', '中国民生银行'),
        (r'^光大', '中国光大银行'),
        (r'^浦发银行?', '上海浦东发展银行'),
        (r'^浦发', '上海浦东发展银行'),
        (r'^广发银行?', '广发银行'),
        (r'^平安银行?', '平安银行'),
        (r'^兴业银行?', '兴业银行')
    ]
    for pat, std_name in prefix_standardize:
        if re.search(pat, s) and not s.startswith(std_name):
            s = re.sub(pat, std_name, s, count=1)
            actions.append(f'规范总行前缀[{std_name}]')
            break

    # 7. 跨地域/跨词重复总行清理（例如“中国农业银行长治市农业银行永泰支行”中跨市县重复出现的“农业银行”）
    for std_bank, aliases in KNOWN_BANK_KEYWORDS:
        if s.startswith(std_bank):
            rest = s[len(std_bank):]
            check_words = [std_bank.replace('中国', ''), std_bank] + aliases
            for w in sorted(check_words, key=len, reverse=True):
                if len(w) >= 2 and w in rest:
                    # 剔除支行部分冗余复现的总行名
                    rest_cleaned = rest.replace(w, '', 1)
                    s = std_bank + rest_cleaned
                    actions.append(f'剔除跨词重复总行[{w}]')
                    break
            break

    # 8. 末尾漏字修补
    if s.endswith('支') and not s.endswith('分支') and not s.endswith('支行'):
        s += '行'
        actions.append('补齐末尾[行]字')
    elif s.endswith('分') and not s.endswith('分行') and len(s) >= 4:
        s += '行'
        actions.append('补齐末尾[行]字')

    # 9. 连续叠字清理
    s = re.sub(r'支行支行', '支行', s)
    s = re.sub(r'分行分行', '分行', s)
    s = re.sub(r'信用社信用社', '信用社', s)

    return s.strip(), actions


def auto_detect_table_structure(df_raw: pd.DataFrame) -> Tuple[pd.DataFrame, Dict[str, Optional[str]], int]:
    """
    智能解析金融结算表格结构：
    1. 自动探测真正表头行（自动跳过报告大标题、合并单元格、顶部空行）
    2. 依据相关文字关键词与列内容特征，动态精准定位：
       - 开户行列 (模糊匹配：开户行/开户支行/银行名称/开户银行/结算银行/支行名称/网点...)
       - 银行卡号列 (模糊匹配：银行卡号/卡号/结算账号/借记卡号/收款账号... + 15-20位卡号内容特征识别)
       - 姓名列 (模糊匹配：姓名/医生姓名/专家姓名/持卡人/收款人/户名...)
    返回: (规范后的DataFrame, 映射字典, 识别到的表头行索引)
    """
    bank_kws = ['开户行', '开户支行', '银行名称', '开户银行', '支行名称', '所属银行', '结算银行', '收款银行', '银行及支行', '开户网点', '支行', '银行']
    card_kws = ['银行卡号', '银行卡', '卡号', '收款卡号', '结算卡号', '结算账号', '银行账号', '收款账号', '借记卡号', '借记卡', '账号', '电签卡号', '电签银行卡号']
    name_kws = ['姓名', '专家姓名', '医生姓名', '持卡人', '收款人', '专家', '医生', '收款人姓名', '户名']
    all_kws = bank_kws + card_kws + name_kws + ['序号', '身份证', '身份证号', '手机号', '电话', '金额', '劳务费', '医院', '职称', '税后']

    # 评估当前 columns 是否直接就是有效表头
    curr_score = sum(1 for c in df_raw.columns if any(k in str(c) for k in all_kws))
    has_unnamed = any('Unnamed' in str(c) for c in df_raw.columns)
    
    best_row_idx = -1  # -1 表示现有的 columns 就是表头
    best_score = curr_score if not has_unnamed else 0
    
    # 扫描前 10 行探测表头
    for r in range(min(10, len(df_raw))):
        row_vals = [str(x).strip() for x in df_raw.iloc[r] if pd.notna(x)]
        score = sum(1 for v in row_vals if any(k in v for k in all_kws))
        if score > best_score:
            best_score = score
            best_row_idx = r
            
    if best_row_idx >= 0:
        raw_row_cols = [str(x).strip() for x in df_raw.iloc[best_row_idx]]
        df_data = df_raw.iloc[best_row_idx+1:].copy().reset_index(drop=True)
        # 排除之前误操作遗留的空列名或重复列名
        seen_cols = {}
        unique_cols = []
        for c in raw_row_cols:
            c_str = str(c).strip()
            if not c_str or c_str.startswith('Unnamed:'):
                unique_cols.append(c_str)
                continue
            if c_str in seen_cols:
                seen_cols[c_str] += 1
                unique_cols.append(f"{c_str}_{seen_cols[c_str]}")
            else:
                seen_cols[c_str] = 0
                unique_cols.append(c_str)
        df_data.columns = unique_cols
        header_excel_line = best_row_idx + 2  # Excel 真实行号（从1起算，且第一行可能为原标题）
    else:
        df_data = df_raw.copy()
        header_excel_line = 1

    # 自动剥离若之前运行残留的质检追加列，防止重复追加
    audit_cols_to_drop = [c for c in df_data.columns if any(k in str(c) for k in ['【清洗后】', '【开户行质检状态】', '【质检核验说明】', '【原始手填备份】', '【规范开户行】'])]
    if audit_cols_to_drop:
        df_data = df_data.drop(columns=audit_cols_to_drop, errors='ignore')

    # 动态抓取列名（过滤纯空列与系统追加列）
    cols = [str(c).strip() for c in df_data.columns if str(c).strip() and not str(c).startswith('Unnamed:') and not str(c).startswith('【')]
    if not cols:
        cols = list(df_data.columns)

    # 1. 动态抓取开户行/银行列
    bank_col = next((c for c in cols if any(k in str(c) for k in ['开户行', '开户支行', '开户银行', '银行名称', '结算银行', '收款银行', '所属银行', '支行名称', '网点名称', '开户网点'])), None)
    if not bank_col:
        bank_col = next((c for c in cols if any(k in str(c) for k in ['银行', '支行', '网点']) and not any(bad in str(c) for bad in ['卡号', '账号', '行号', '代码', '卡', '账'])), None)
    if not bank_col:
        # 数据内容特征扫描：包含支行/分行/营业室/信用社
        for c in cols:
            samples = [str(x) for x in df_data[c].dropna().head(10)]
            if any('支行' in s or '银行' in s or '分行' in s or '营业' in s or '信用社' in s for s in samples):
                bank_col = c
                break

    # 2. 动态抓取银行卡号列
    card_col = next((c for c in cols if any(k in str(c) for k in ['银行卡号', '银行卡', '卡号', '收款卡号', '结算卡号', '借记卡号', '结算账号', '银行账号', '收款账号', '借记卡', '账号', '电签银行卡号']) and not any(bad in str(c) for bad in ['开户行', '支行', '行名'])), None)
    if not card_col:
        # 数据内容特征扫描：15-20位连续纯数字
        for c in cols:
            samples = [re.sub(r'[\s\.\-]', '', str(x)) for x in df_data[c].dropna().head(10)]
            digit_cnt = sum(1 for s in samples if re.match(r'^\d{15,20}$', s))
            if digit_cnt >= 2:
                card_col = c
                break

    # 3. 动态抓取姓名列
    name_col = next((c for c in cols if any(k in str(c) for k in ['姓名', '医生姓名', '专家姓名', '持卡人', '收款人', '户名', '收款人姓名', '专家', '医生']) and '卡' not in str(c)), None)

    return df_data, {'bank_col': bank_col, 'card_col': card_col, 'name_col': name_col, 'header_row': header_excel_line}


def clean_single_bank_record(
    raw_bank: str,
    card_no: Optional[str] = None,
    doctor_name: Optional[str] = None
) -> Dict[str, str]:
    """
    单条记录清洗质检引擎
    """
    raw_str = str(raw_bank).strip() if pd.notna(raw_bank) else ""
    if raw_str in ["nan", "None", "null", "NULL"]:
        raw_str = ""

    result = {
        "raw_bank": raw_str,
        "cleaned_bank": raw_str,
        "status": "✅ 正常",
        "detail": "格式合规",
        "is_repaired": False,
        "repair_note": ""
    }

    if not raw_str:
        result["status"] = "❌ 缺失"
        result["detail"] = "开户行信息完全为空"
        return result

    # 1. 检查是否误填为医生姓名
    if doctor_name and str(doctor_name).strip() and raw_str == str(doctor_name).strip():
        result["status"] = "❌ 误填姓名"
        result["detail"] = f"误填为了医生本人姓名【{doctor_name}】"
        result["cleaned_bank"] = ""
        return result

    # 2. 执行强力去重、清理前缀与格式规范化
    cleaned_text, dedup_actions = advanced_dedup_and_clean_bank(raw_str)
    curr = cleaned_text
    if dedup_actions:
        result["is_repaired"] = True
        result["repair_note"] += "; ".join(dedup_actions) + "; "

    # 3. 银行总行识别
    inferred_bank = identify_bank_from_card(card_no) if card_no else None

    has_bank_name = False
    found_main_bank = ""
    for full_b, aliases in KNOWN_BANK_KEYWORDS:
        if full_b in curr:
            has_bank_name = True
            found_main_bank = full_b
            break
        for a in aliases:
            if a in curr:
                has_bank_name = True
                found_main_bank = full_b
                break
        if has_bank_name:
            break

    if not has_bank_name:
        if "银行" in curr or "信用社" in curr or "合作社" in curr:
            has_bank_name = True

    # 4. 缺少总行（只有支行）的智能补全
    if not has_bank_name:
        if inferred_bank:
            # 奇迹补全：有卡号，卡号 BIN 码自动反查总行并拼接
            curr = f"{inferred_bank}{curr}"
            result["is_repaired"] = True
            result["repair_note"] += f"卡号BIN反查补齐总行【{inferred_bank}】; "
            has_bank_name = True
            found_main_bank = inferred_bank
        else:
            # 无卡号时：根据专属行业网点特征智能启发推断
            if any(kw in curr for kw in ["住房支行", "房建支行", "房产支行", "公积金支行"]):
                curr = f"中国建设银行{curr}"
                result["is_repaired"] = True
                result["repair_note"] += "依据【住房支行】行业网点特征智能补全总行【中国建设银行】; "
                has_bank_name = True
                found_main_bank = "中国建设银行"
                result["status"] = "✅ 正常(智能推断)"
                result["detail"] = "依据行业网点特征补齐建行，建议核实"
            else:
                result["status"] = "⚠️ 缺少银行"
                result["detail"] = "缺少银行总行名称(仅填支行)，请提供卡号自动补齐或人工补全"

    # 5. 支行/分行等网点层级质检
    has_branch_level = any(bk in curr for bk in BRANCH_KEYWORDS)
    if not has_branch_level:
        if result["status"].startswith("✅"):
            result["status"] = "⚠️ 缺少网点"
            result["detail"] = "缺少支行/分行/营业室等具体网点层级，转账极易退票"

    # 6. 卡号与填写银行冲突核验
    if inferred_bank and found_main_bank:
        if (inferred_bank not in found_main_bank) and (found_main_bank not in inferred_bank):
            is_mismatch = True
            for full_b, aliases in KNOWN_BANK_KEYWORDS:
                if (full_b == inferred_bank and found_main_bank in aliases) or \
                   (full_b == found_main_bank and inferred_bank in aliases):
                    is_mismatch = False
                    break
            if is_mismatch:
                result["status"] = "⚠️ 卡号与银行冲突"
                result["detail"] = f"卡号前缀识别为【{inferred_bank}】，但填写的为【{found_main_bank}】"

    # 7. 内容过短检测
    if len(curr) < 4:
        result["status"] = "❌ 内容过短"
        result["detail"] = "开户行名称过短，无法正常汇款"

    result["cleaned_bank"] = curr
    if result["status"] == "✅ 正常" and result["is_repaired"]:
        result["detail"] = "已自动修复并格式达标"

    return result


def clean_dataframe_banks(
    df: pd.DataFrame,
    bank_col: str,
    card_col: Optional[str] = None,
    name_col: Optional[str] = None,
    replace_original: bool = True
) -> Tuple[pd.DataFrame, Dict[str, int]]:
    """
    批量清洗 DataFrame：
    - 若 replace_original=True：直接在原【开户行】列生效替换，并在其右侧紧邻插入【原始手填备份】开户行；
    - 在表末尾追加【开户行质检状态】与【质检核验说明】；
    - 确保财务打开 Excel 一眼看到清洗干净的开户行列，原数据留档备查。
    """
    df_out = df.copy()

    cleaned_banks = []
    statuses = []
    details = []
    
    stats = {
        "total": len(df_out),
        "normal": 0,
        "repaired": 0,
        "warning": 0,
        "error": 0
    }

    # 第一轮：全量清洗与初步质检，并提取本批次已合规网点知识库
    first_pass_results = []
    valid_batch_branches = set()

    for idx, row in df_out.iterrows():
        raw_b = row.get(bank_col, "")
        card_v = row.get(card_col, "") if card_col else None
        name_v = row.get(name_col, "") if name_col else None

        res = clean_single_bank_record(raw_b, card_v, name_v)
        first_pass_results.append(res)
        if res["status"].startswith("✅"):
            valid_batch_branches.add(res["cleaned_bank"])

    # 第二轮：基于批次知识库与地名特征上下文，对漏写“支行”的记录二次智能对齐
    for res in first_pass_results:
        curr_b = res["cleaned_bank"]
        if res["status"].startswith("⚠️") and "缺少网点" in res["status"]:
            candidate_branch = curr_b + "支行"
            # 1. 优先命中同批次已有的合规支行（如“中国银行太原杏花岭”命中同批的“中国银行太原杏花岭支行”）
            if candidate_branch in valid_batch_branches:
                res["cleaned_bank"] = candidate_branch
                res["status"] = "✅ 正常"
                res["detail"] = "已自动修复并格式达标"
                res["repair_note"] += "补齐末尾[支行]字 (依据同批已合规网点对齐); "
                res["is_repaired"] = True
            # 2. 网点以行政区/街道/地名结尾（如杏花岭、建设路、迎泽），但漏写支行
            elif re.search(r'(?:区|县|镇|街|路|道|岭|桥|门|城|园|巷|矿|湾|坪|港|河|洲|湖|岛|堡|铺|庄)$', curr_b):
                res["cleaned_bank"] = candidate_branch
                res["status"] = "✅ 正常"
                res["detail"] = "已自动修复并格式达标"
                res["repair_note"] += "补齐末尾[支行]字; "
                res["is_repaired"] = True

        cleaned_banks.append(res["cleaned_bank"])
        statuses.append(res["status"])
        
        full_detail = res["detail"]
        if res["repair_note"]:
            full_detail += f" ({res['repair_note'].rstrip('; ')})"
        details.append(full_detail)

        if res["is_repaired"]:
            stats["repaired"] += 1

        if res["status"].startswith("✅"):
            stats["normal"] += 1
        elif res["status"].startswith("⚠️"):
            stats["warning"] += 1
        else:
            stats["error"] += 1

    if replace_original:
        # 1. 备份原手填列
        orig_series = df_out[bank_col].copy()
        # 2. 原列直接覆盖为清洗后规范全称
        df_out[bank_col] = cleaned_banks
        # 3. 在原列后插入备份列
        col_loc = df_out.columns.get_loc(bank_col) + 1
        df_out.insert(col_loc, "【原始手填备份】开户行", orig_series)
        # 4. 表末尾追加状态与说明
        df_out["【开户行质检状态】"] = statuses
        df_out["【质检核验说明】"] = details
    else:
        df_out["【清洗后】开户行规范全称"] = cleaned_banks
        df_out["【开户行质检状态】"] = statuses
        df_out["【质检核验说明】"] = details

    return df_out, stats


def export_styled_excel(df: pd.DataFrame, output_path: str):
    """
    导出高保真带样式 Excel（强化文本保护与浅红高亮）
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    # 确保输出文件路径为标准 .xlsx
    output_path = str(Path(output_path).with_suffix(".xlsx"))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "开户行清洗质检明细"

    # 样式定义
    font_header = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")

    font_data = Font(name="微软雅黑", size=10)
    font_status_ok = Font(name="微软雅黑", size=10, bold=True, color="15803D")
    font_status_warn = Font(name="微软雅黑", size=10, bold=True, color="B45309")
    font_status_err = Font(name="微软雅黑", size=10, bold=True, color="B91C1C")

    fill_warn_row = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")
    fill_err_row = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    headers = list(df.columns)
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

    status_col_pos = headers.index("【开户行质检状态】") if "【开户行质检状态】" in headers else None

    for r_idx in range(len(df)):
        excel_row_num = r_idx + 2
        status_val = str(df.iloc[r_idx, status_col_pos]) if status_col_pos is not None else ""

        for c_idx in range(1, len(headers) + 1):
            col_name = headers[c_idx - 1]
            val = df.iloc[r_idx, c_idx - 1]
            if pd.isna(val):
                val = ""
            
            # 卡号、手机号、身份证号防科学计数法
            col_str = str(col_name)
            is_sensitive_num = any(k in col_str for k in ["卡号", "账号", "手机", "电话", "身份证", "编码", "编号"])
            if is_sensitive_num and val != "":
                val = clean_card_number(val)

            cell = ws.cell(row=excel_row_num, column=c_idx, value=val)
            cell.font = font_data
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", horizontal="left")

            if is_sensitive_num:
                cell.number_format = "@"

            # 异常行浅黄/浅红高亮
            if "⚠️" in status_val:
                cell.fill = fill_warn_row
            elif "❌" in status_val:
                cell.fill = fill_err_row

        if status_col_pos is not None:
            scell = ws.cell(row=excel_row_num, column=status_col_pos + 1)
            scell.alignment = Alignment(horizontal="center", vertical="center")
            if "✅" in status_val:
                scell.font = font_status_ok
            elif "⚠️" in status_val:
                scell.font = font_status_warn
            else:
                scell.font = font_status_err

        ws.row_dimensions[excel_row_num].height = 22

    ws.freeze_panes = "A2"

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            length = sum(2 if ord(c) > 127 else 1 for c in val_str)
            if length > max_len:
                max_len = length
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 48)

    wb.save(output_path)
    return output_path
